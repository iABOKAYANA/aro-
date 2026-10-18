"""Extract the owner-supplied SODIC XLSX into a PII-free, mapped import JSON."""
import collections
import datetime
import json
import re
import sys

import openpyxl

ROUTE = {
    ("464 Acre", "Karmell"): {"target": "402-karmell-new-zayed"},
    ("464 Acre", "VYE"): {"target": "336-vye-sodic", "phase": "VYE"},
    ("464 Acre", "Commercial"): {"new": {"slug": "464-acre-commercial-sodic", "name": "464 Acre Commercial", "nameArabic": "464 Acre التجاري"}},
    ("Caesar", "Extension 1"): {"target": "1783-caesar-extension-1"},
    ("Caesar", "Extension 2"): {"target": "1696-caesar-extension-2"},
    ("EASTVALE", "Phase 1"): {"target": "1846-sodic-mostakbal-city"},
    ("Forty West", "Forty West 1"): {"target": "26-forty-west-phase-i"},
    ("JUNE", "JUNE I"): {"target": "510-june", "phase": "JUNE I"},
    ("OGAMI", "Botanica Town"): {"target": "2176-botanica-town---ogami"},
    ("OGAMI", "Nobu Parcel"): {"target": "2612-nobu-residence"},
    ("OGAMI", "OGM P1"): {"target": "1237-ogami-ras-el-hekma", "phase": "OGM P1"},
    ("SIX WEST", "SIX WEST 1"): {"target": "233-six-west"},
    ("SODIC EAST", "SODIC EAST 1"): {"target": "1617-orchid"},
    ("SODIC EAST", "SODIC EAST 2"): {"target": "176-sodic-east", "phase": "SODIC EAST 2"},
    ("SODIC EAST", "SODIC EAST 2 MF"): {"target": "2726-the-lakes-sodic"},
    ("SODIC EAST", "SODIC EAST 2-"): {"target": "176-sodic-east", "phase": "SODIC EAST 2-"},
    ("THE ESTATES", "THE ESTATES 3"): {"target": "310-the-estates", "phase": "THE ESTATES 3"},
    ("THE ESTATES", "THE ESTATES 4"): {"target": "310-the-estates", "phase": "THE ESTATES 4"},
    ("THE ESTATES Residences", "THE ESTATES Residences II (MF)"): {"target": "664-the-estates-residence", "phase": "THE ESTATES Residences II (MF)"},
    ("THE ESTATES Residences", "THE ESTATES Residences III (MF)"): {"target": "664-the-estates-residence", "phase": "THE ESTATES Residences III (MF)"},
    ("THE ESTATES Residences", "The Estates Residences 1 Extension"): {"target": "664-the-estates-residence", "phase": "The Estates Residences 1 Extension"},
    ("THE ESTATES Residences", "The Estates Residences III (SF)"): {"target": "664-the-estates-residence", "phase": "The Estates Residences III (SF)"},
    ("THE PORTAL", "THE PORTAL 1"): {"target": "638-westown-the-portal"},
    ("The Strip II", "The Strip 2"): {"target": "28-the-strip-[-commercial-]"},
    ("V Editions", "V Editions"): {"new": {"slug": "v-editions-sodic", "name": "V Editions", "nameArabic": "ڤي إديشنز"}},
    ("Villette", "SKY CONDOS 1"): {"target": "654-sky-condos", "phase": "SKY CONDOS 1"},
    ("Villette", "SKY CONDOS 2"): {"target": "654-sky-condos", "phase": "SKY CONDOS 2"},
    ("Villette", "SKY CONDOS 3"): {"target": "180-sky-condos-phase-3,-villette"},
    ("Villette", "V Residence"): {"target": "260-v-residences-by-villette"},
    ("Westown Medical Centre", "WMC 1"): {"target": "1052-westown-medical-center"},
}

TYPE_MAP = {
    "Apartment": "APARTMENT", "Duplex": "DUPLEX", "Villa": "VILLA",
    "Town House": "TOWNHOUSE", "Twin House": "VILLA", "Office": "ADMINISTRATIVE",
    "House": "VILLA", "Loft": "APARTMENT", "Chalet": "CHALET",
    "Serviced Apartments": "APARTMENT", "Branded Residences": "APARTMENT",
    "Shops": "COMMERCIAL", "Clinic": "MEDICAL",
}
TYPE_AR = {
    "APARTMENT": "شقة", "DUPLEX": "دوبلكس", "VILLA": "فيلا", "TOWNHOUSE": "تاون هاوس",
    "ADMINISTRATIVE": "إداري", "CHALET": "شاليه", "COMMERCIAL": "محل تجاري", "MEDICAL": "عيادة",
}
MONTHS = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")


def number(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def delivery(value):
    if isinstance(value, (datetime.date, datetime.datetime)):
        return f"Estimated delivery: {value.day:02d} {MONTHS[value.month - 1]} {value.year}"
    text = str(value or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            d = datetime.datetime.strptime(text, fmt)
            return f"Estimated delivery: {d.day:02d} {MONTHS[d.month - 1]} {d.year}"
        except ValueError:
            pass
    return f"Estimated delivery: {text}" if text else None


def main(source, output):
    ws = openpyxl.load_workbook(source, data_only=True, read_only=True).active
    headers = []
    for value in next(ws.iter_rows(min_row=12, max_row=12, values_only=True)):
        text = str(value).strip() if value is not None else ""
        headers.append(re.sub(r"\s*↑\s*$", "", text).strip())

    rows, unmapped, invalid = [], collections.Counter(), []
    project, unit_type = "", ""
    for values in ws.iter_rows(min_row=13, values_only=True):
        d = dict(zip(headers, values))
        if str(d.get("Project Name") or "").strip():
            project = str(d["Project Name"]).strip()
        if str(d.get("Unit Type") or "").strip():
            unit_type = str(d["Unit Type"]).strip()
        code = str(d.get("Unit Name") or "").strip()
        if not code:
            continue

        stage = str(d.get("Phase: Phase Name") or "").strip()
        route = ROUTE.get((project, stage))
        if not route:
            unmapped[(project, stage)] += 1
            continue
        kind = TYPE_MAP.get(unit_type)
        # This one available row has no type or area in the source; the source name
        # says "App Unit", so classify as apartment and retain the missing area.
        unknown_lakes_app = code == "SODIC Community App Unit - Lakes - 1"
        if not kind and unknown_lakes_app:
            kind, unit_type = "APARTMENT", "Community App Unit (source type blank)"
        price = number(d.get("Final Total Unit Price"))
        gross, sellable = number(d.get("Unit Gross Area")), number(d.get("Sellable Unit Area"))
        area = gross if gross and gross > 0 else (sellable if sellable and sellable > 0 else (0 if unknown_lakes_app else None))
        if not kind or not price or price <= 0 or area is None:
            invalid.append(code)
            continue

        rooms = number(d.get("Number of rooms"))
        bedrooms = int(rooms) if rooms is not None and rooms >= 0 and rooms.is_integer() else None
        garden, roof = number(d.get("Garden Area")), number(d.get("Open Roof Deck"))
        garden = garden if garden and garden > 0 else None
        roof = roof if roof and roof > 0 else None
        floor_text = str(d.get("Floor") or "").strip()
        floor = int(floor_text) if floor_text.isdigit() else (0 if floor_text == "G" else None)
        display_type = unit_type if not unknown_lakes_app else "Community App Unit - Lakes"
        en = [display_type]
        ar = ["وحدة Community App في Lakes" if unknown_lakes_app else TYPE_AR[kind]]
        if bedrooms is not None and bedrooms > 0:
            en.append(f"{bedrooms}BR")
            ar.append("غرفة" if bedrooms == 1 else "غرفتين" if bedrooms == 2 else f"{bedrooms} غرف")
        if area > 0:
            en.append(f"{area:g} sqm")
            ar.append(f"{area:g} م²")
        else:
            en.append("area not stated")
            ar.append("المساحة غير مذكورة")
        if floor_text and floor_text not in ("0", "G") and not floor_text.isdigit():
            en.append(f"Floor {floor_text}")
            ar.append(f"الدور {floor_text}")

        route_data = dict(route)
        target = route_data.get("target") or route_data["new"]["slug"]
        rows.append({
            "sourceProject": project, "sourceStage": stage, "target": target,
            "new": route_data.get("new"), "phase": route_data.get("phase"),
            "unitNumber": code, "unitType": kind, "nameEn": " · ".join(en), "nameAr": " · ".join(ar),
            "bedrooms": bedrooms, "area": area, "gardenArea": garden, "roofArea": roof,
            "price": price, "finishingType": str(d.get("Finishing Type") or "").strip() or None,
            "deliveryTerms": delivery(d.get("Estimated Delivery Date")), "floor": floor,
        })

    keys = [(r["target"], r["unitNumber"]) for r in rows]
    print("rows", len(rows), "unmapped", dict(unmapped), "invalid", invalid)
    print("duplicate target/unit keys", len(keys) - len(set(keys)))
    print("targets", dict(sorted(collections.Counter(r["target"] for r in rows).items())))
    if len(rows) != 1300 or unmapped or invalid or len(keys) != len(set(keys)):
        raise SystemExit("Validation failed; refusing to write the import data")
    new_projects = {r["new"]["slug"]: r["new"] for r in rows if r["new"]}
    with open(output, "w", encoding="utf-8") as f:
        json.dump({"sourceDate": "2026-10-15", "newProjects": new_projects, "rows": rows}, f, ensure_ascii=False)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: prepare-sodic-availability.py INPUT.xlsx OUTPUT.json")
    main(sys.argv[1], sys.argv[2])
